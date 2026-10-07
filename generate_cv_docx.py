import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn, nsdecls
from docx.oxml import parse_xml

# Color constants
COLOR_NAVY = RGBColor(30, 58, 138)       # #1E3A8A - Primary accent
COLOR_DARK = RGBColor(15, 23, 42)        # #0F172A - Body text
COLOR_MUTED = RGBColor(71, 85, 105)      # #475569 - Secondary text
HEX_NAVY = "1E3A8A"
HEX_BORDER = "CBD5E1"                    # #CBD5E1 - Light border
HEX_BG_LIGHT = "F8FAFC"                  # #F8FAFC - Soft panel bg

def set_cell_margins(cell, top=100, bottom=100, left=140, right=140):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('w:top', top), ('w:bottom', bottom), ('w:left', left), ('w:right', right)]:
        node = OxmlElement(m)
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def add_hyperlink(paragraph, url, text, color_hex="1E3A8A", bold=False):
    part = paragraph.part
    r_id = part.relate_to(url, docx.opc.constants.RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    hyperlink = OxmlElement('w:hyperlink')
    hyperlink.set(qn('r:id'), r_id)
    new_run = OxmlElement('w:r')
    rPr = OxmlElement('w:rPr')
    if color_hex:
        c = OxmlElement('w:color')
        c.set(qn('w:val'), color_hex)
        rPr.append(c)
    u = OxmlElement('w:u')
    u.set(qn('w:val'), 'single')
    rPr.append(u)
    if bold:
        b = OxmlElement('w:b')
        rPr.append(b)
    new_run.append(rPr)
    new_run.text = text
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)

def add_section_heading(doc, title):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    
    run = p.add_run(title.upper())
    run.font.name = "Calibri"
    run.font.size = Pt(11.5)
    run.font.bold = True
    run.font.color.rgb = COLOR_NAVY

    # Add bottom border under heading (clean 1px hairline)
    pBdr = parse_xml(r'<w:pBdr xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
                     r'<w:bottom w:val="single" w:sz="6" w:space="3" w:color="1E3A8A"/>'
                     r'</w:pBdr>')
    p._p.get_or_add_pPr().append(pBdr)

def build_cv_sk():
    doc = docx.Document()
    
    # Page Setup (A4, 1.8 cm margins)
    section = doc.sections[0]
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)

    # 1. HEADER
    p_name = doc.add_paragraph()
    p_name.paragraph_format.space_before = Pt(0)
    p_name.paragraph_format.space_after = Pt(2)
    run_name = p_name.add_run("Ing. Ruslana Buinytska")
    run_name.font.name = "Calibri"
    run_name.font.size = Pt(22)
    run_name.font.bold = True
    run_name.font.color.rgb = COLOR_NAVY

    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(6)
    run_title = p_title.add_run("IT Analytička (Systémová & Biznis Analýza)  •  IT Project Manager")
    run_title.font.name = "Calibri"
    run_title.font.size = Pt(12)
    run_title.font.bold = True
    run_title.font.color.rgb = COLOR_DARK

    # Contact line 1: Location, Phone, Email
    p_contact1 = doc.add_paragraph()
    p_contact1.paragraph_format.space_before = Pt(0)
    p_contact1.paragraph_format.space_after = Pt(2)
    
    r = p_contact1.add_run("Bratislava, Slovensko  |  Tel: ")
    r.font.size = Pt(9.5); r.font.color.rgb = COLOR_MUTED
    r_tel = p_contact1.add_run("+421 951 698 832")
    r_tel.font.size = Pt(9.5); r_tel.font.bold = True; r_tel.font.color.rgb = COLOR_DARK
    
    r2 = p_contact1.add_run("  |  Email: ")
    r2.font.size = Pt(9.5); r2.font.color.rgb = COLOR_MUTED
    add_hyperlink(p_contact1, "mailto:rusbuin.oct@gmail.com", "rusbuin.oct@gmail.com", HEX_NAVY)

    # Contact line 2: Web Portfolio, LinkedIn, GitHub
    p_contact2 = doc.add_paragraph()
    p_contact2.paragraph_format.space_before = Pt(0)
    p_contact2.paragraph_format.space_after = Pt(10)
    
    r3 = p_contact2.add_run("Web Portfólio: ")
    r3.font.size = Pt(9.5); r3.font.bold = True; r3.font.color.rgb = COLOR_NAVY
    add_hyperlink(p_contact2, "https://rbuinytska.github.io", "rbuinytska.github.io", HEX_NAVY, bold=True)
    
    r4 = p_contact2.add_run("  |  LinkedIn: ")
    r4.font.size = Pt(9.5); r4.font.color.rgb = COLOR_MUTED
    add_hyperlink(p_contact2, "https://linkedin.com/in/ruslana-buinytska-188602395", "linkedin.com/in/ruslana-buinytska", HEX_NAVY)
    
    r5 = p_contact2.add_run("  |  GitHub: ")
    r5.font.size = Pt(9.5); r5.font.color.rgb = COLOR_MUTED
    add_hyperlink(p_contact2, "https://github.com/rbuinytska", "github.com/rbuinytska", HEX_NAVY)

    # 2. PROFIL
    add_section_heading(doc, "Profesijný Profil")
    p_profile = doc.add_paragraph()
    p_profile.paragraph_format.space_before = Pt(4)
    p_profile.paragraph_format.space_after = Pt(8)
    p_profile.paragraph_format.line_spacing = 1.15
    run_prof = p_profile.add_run(
        "Inžinierka Informačného manažmentu (EUBA) so špecializáciou na IT systémovú analýzu, riadenie IT projektov "
        "a dátovú architektúru. Vďaka praktickej skúsenosti zo softvérového vývoja rozumiem systémovej architektúre "
        "do hĺbky, čo mi umožňuje tvoriť precízne funkčné a technické špecifikácie bez logických medzier, navrhovať "
        "integračné rozhrania (REST API, JSON schémy), riadiť agilné tímy (certifikácia IPMA, SCRUM) a byť "
        "rovnocenným partnerom pre vývojárov aj manažment."
    )
    run_prof.font.size = Pt(10)
    run_prof.font.color.rgb = COLOR_DARK

    # 3. KĽÚČOVÉ KOMPETENCIE
    add_section_heading(doc, "Kľúčové Kompetencie & Metodiky")
    
    skills_data = [
        ("IT Analýza & Modelovanie:", "UML (Use Case, Activity, Sequence diagramy), BPMN procesy, Enterprise Architect (EA), Draw.io, funkčné a technické špecifikácie, zber a dekompozícia biznis požiadaviek, návrh API kontraktov (REST, JSON), User Stories, akceptačné kritériá."),
        ("Projektové Riadenie:", "Medzinárodná certifikácia IPMA Project Management Fundamentals (2025), agilné rámce SCRUM & Kanban, prioritizácia sprint backlogu, riadenie rizík, harmonogram dodávky, princípy PRINCE2 a ITIL."),
        ("Dáta & BI Analytika:", "Relačné databázy, pokročilé Oracle SQL, dátové modelovanie (Star Schema), ETL transformácie tokov dát, interaktívne dashboardy v Power BI (DAX) a Tableau, dátová analýza v Pythone."),
        ("Technický Vývoj & AI:", "REST API architektúra, JSON/XML, Kotlin (KMP), Java, Android, Git, AI-Driven Development (LLM integrácia, tvorba a orchestrácia AI agentov).")
    ]
    
    for cat_title, cat_desc in skills_data:
        p_sk = doc.add_paragraph(style='List Bullet')
        p_sk.paragraph_format.space_before = Pt(2)
        p_sk.paragraph_format.space_after = Pt(2)
        p_sk.paragraph_format.line_spacing = 1.1
        r_t = p_sk.add_run(cat_title + " ")
        r_t.font.bold = True
        r_t.font.size = Pt(9.5)
        r_t.font.color.rgb = COLOR_NAVY
        r_d = p_sk.add_run(cat_desc)
        r_d.font.size = Pt(9.5)
        r_d.font.color.rgb = COLOR_DARK

    # 4. PRACOVNÉ SKÚSENOSTI
    add_section_heading(doc, "Pracovné Skúsenosti")

    # Item: WBPO
    p_wbpo_h = doc.add_paragraph()
    p_wbpo_h.paragraph_format.space_before = Pt(6)
    p_wbpo_h.paragraph_format.space_after = Pt(2)
    p_wbpo_h.paragraph_format.keep_with_next = True
    r_role = p_wbpo_h.add_run("IT Analýza & Softvérový Vývoj")
    r_role.font.bold = True; r_role.font.size = Pt(11); r_role.font.color.rgb = COLOR_DARK
    r_comp = p_wbpo_h.add_run("  |  WBPO s.r.o. (Bratislava, Slovensko)")
    r_comp.font.bold = True; r_comp.font.size = Pt(10.5); r_comp.font.color.rgb = COLOR_NAVY
    
    p_wbpo_d = doc.add_paragraph()
    p_wbpo_d.paragraph_format.space_before = Pt(0)
    p_wbpo_d.paragraph_format.space_after = Pt(3)
    p_wbpo_d.paragraph_format.keep_with_next = True
    r_date = p_wbpo_d.add_run("2025 – súčasnosť  •  Aktuálna pozícia")
    r_date.font.size = Pt(9); r_date.font.italic = True; r_date.font.color.rgb = COLOR_MUTED

    wbpo_bullets = [
        "Samostatná dekompozícia biznisových požiadaviek stakeholderov do detailnej technickej analýzy a funkčných špecifikácií aplikačnej logiky pre vývojový tím.",
        "Návrh systémových integračných rozhraní (REST API kontrakty, dátové JSON schémy) a aktívna koordinácia technickej realizácie s backend architektmi a vývojármi.",
        "Definícia akceptačných kritérií, tvorba detailných testovacích scenárov, identifikácia a ošetrenie okrajových stavov (edge cases) a výnimiek.",
        "Práca v agilnom SCRUM tíme, účasť na sprint planningoch, refinemantoch zadaní a prepájanie biznisových očakávaní s technickou realitou."
    ]
    for b in wbpo_bullets:
        p_b = doc.add_paragraph(style='List Bullet')
        p_b.paragraph_format.space_before = Pt(1)
        p_b.paragraph_format.space_after = Pt(2)
        p_b.paragraph_format.line_spacing = 1.1
        r_b = p_b.add_run(b)
        r_b.font.size = Pt(9.5)
        r_b.font.color.rgb = COLOR_DARK

    p_wbpo_t = doc.add_paragraph()
    p_wbpo_t.paragraph_format.space_before = Pt(2)
    p_wbpo_t.paragraph_format.space_after = Pt(6)
    r_tl = p_wbpo_t.add_run("Technológie & Nástroje: ")
    r_tl.font.bold = True; r_tl.font.size = Pt(9); r_tl.font.color.rgb = COLOR_MUTED
    r_tv = p_wbpo_t.add_run("IT Systémová Analýza, REST API, JSON, UML, Agile / SCRUM, Technické špecifikácie, Kotlin (KMP), Git")
    r_tv.font.size = Pt(9); r_tv.font.color.rgb = COLOR_MUTED

    # 5. VYBRANÉ PROJEKTY
    add_section_heading(doc, "Vybrané Projekty & Analytické Portfólio")

    # Project 1: ESG
    p_esg_h = doc.add_paragraph()
    p_esg_h.paragraph_format.space_before = Pt(6)
    p_esg_h.paragraph_format.space_after = Pt(2)
    p_esg_h.paragraph_format.keep_with_next = True
    r_p1 = p_esg_h.add_run("ESG Data Analýza a Vizualizácia  ")
    r_p1.font.bold = True; r_p1.font.size = Pt(10.5); r_p1.font.color.rgb = COLOR_DARK
    r_p1s = p_esg_h.add_run("|  Diplomová práca (Ing.), EUBA  |  ")
    r_p1s.font.size = Pt(9.5); r_p1s.font.color.rgb = COLOR_MUTED
    add_hyperlink(p_esg_h, "https://github.com/RusBuin/esg-analysis-powerbi", "github.com/RusBuin/esg-analysis-powerbi", HEX_NAVY)

    esg_bullets = [
        "End-to-end analytický projekt: zber, validácia, čistenie a ETL transformácia nestruktúrovaných dát environmentálneho a sociálneho riadenia (ESG).",
        "Návrh viacrozmerného dátového modelu (Star Schema) a implementácia analytických biznis metrík a kalkulácií v jazyku DAX.",
        "Tvorba interaktívnych reportingových dashboardov v Power BI pre podporu strategického manažérskeho rozhodovania."
    ]
    for b in esg_bullets:
        p_b = doc.add_paragraph(style='List Bullet')
        p_b.paragraph_format.space_before = Pt(1)
        p_b.paragraph_format.space_after = Pt(2)
        r_b = p_b.add_run(b)
        r_b.font.size = Pt(9.5); r_b.font.color.rgb = COLOR_DARK

    # Project 2: AI / ML
    p_ml_h = doc.add_paragraph()
    p_ml_h.paragraph_format.space_before = Pt(4)
    p_ml_h.paragraph_format.space_after = Pt(2)
    p_ml_h.paragraph_format.keep_with_next = True
    r_p2 = p_ml_h.add_run("Trénovanie AI Modelu a Dátová Analýza  ")
    r_p2.font.bold = True; r_p2.font.size = Pt(10.5); r_p2.font.color.rgb = COLOR_DARK
    r_p2s = p_ml_h.add_run("|  Bakalárska práca (Bc.), TUKE  |  ")
    r_p2s.font.size = Pt(9.5); r_p2s.font.color.rgb = COLOR_MUTED
    add_hyperlink(p_ml_h, "https://github.com/RusBuin/bakalarska_praca.git", "github.com/RusBuin/bakalarska_praca", HEX_NAVY)

    ml_bullets = [
        "Analýza rozsiahlych datasetov, štatistické predspracovanie dát a natrénovanie predikčného modelu strojového učenia v Pythone.",
        "Optimalizácia algoritmov, vyhodnocovanie presnosti modelov a aplikácia logického riešenia analytických výziev."
    ]
    for b in ml_bullets:
        p_b = doc.add_paragraph(style='List Bullet')
        p_b.paragraph_format.space_before = Pt(1)
        p_b.paragraph_format.space_after = Pt(2)
        r_b = p_b.add_run(b)
        r_b.font.size = Pt(9.5); r_b.font.color.rgb = COLOR_DARK

    # 6. VZDELANIE
    add_section_heading(doc, "Vzdelanie")
    
    edu_list = [
        ("2026 – súčasnosť", "Data Science v ekonómii (PhD.)", "Ekonomická univerzita v Bratislave (EUBA)", "Doktorandské štúdium a výskum zameraný na pokročilé dátové metódy a AI v ekonomických systémoch."),
        ("2024 – 2026", "Informačný manažment (Ing.)", "Ekonomická univerzita v Bratislave (EUBA)", "Inžinierske štúdium: riadenie IT procesov, podnikové informačné systémy, IT systémová analýza a projektový manažment."),
        ("2022 – 2024", "Hospodárska informatika (Bc.)", "Technická univerzita v Košiciach (TUKE)", "Bakalárske štúdium: hospodárska informatika, softvérové inžinierstvo, relačné databázy a dátové štruktúry.")
    ]
    for yr, deg, schl, desc in edu_list:
        p_e = doc.add_paragraph()
        p_e.paragraph_format.space_before = Pt(3)
        p_e.paragraph_format.space_after = Pt(2)
        r_deg = p_e.add_run(f"{deg}  |  {schl}")
        r_deg.font.bold = True; r_deg.font.size = Pt(10); r_deg.font.color.rgb = COLOR_DARK
        r_yr = p_e.add_run(f"  ({yr})\n")
        r_yr.font.size = Pt(9); r_yr.font.color.rgb = COLOR_MUTED
        r_dc = p_e.add_run(desc)
        r_dc.font.size = Pt(9.5); r_dc.font.color.rgb = COLOR_MUTED

    # 7. CERTIFIKÁTY & ŠKOLENIA
    add_section_heading(doc, "Certifikáty & Odborné Kurzy")
    certs = [
        ("2025", "Certifikát IPMA Project Management Fundamentals", "IPMA Slovensko – Medzinárodná certifikácia riadenia projektov, plánovania a kontroly."),
        ("2025", "Sustainable Project Management #HACKATHON2025", "Praktické riešenie udržateľných výziev s využitím agilných projektových prístupov a tímového vedenia."),
        ("2022", "Agile workshop „LEGO SCRUM & TUKE“", "Technická univerzita v Košiciach – Praktická znalosť agilných ceremónií a SCRUM frameworku."),
        ("2022", "Kurz „Skills for Success“ (AmCham Slovakia)", "Americká obchodná komora – Intenzívny rozvoj manažérskych zručností, vyjednávania a biznis komunikácie.")
    ]
    for yr, title, desc in certs:
        p_c = doc.add_paragraph(style='List Bullet')
        p_c.paragraph_format.space_before = Pt(1)
        p_c.paragraph_format.space_after = Pt(2)
        r_ct = p_c.add_run(f"{title} ({yr}): ")
        r_ct.font.bold = True; r_ct.font.size = Pt(9.5); r_ct.font.color.rgb = COLOR_DARK
        r_cd = p_c.add_run(desc)
        r_cd.font.size = Pt(9.5); r_cd.font.color.rgb = COLOR_MUTED

    # 8. JAZYKY
    add_section_heading(doc, "Jazykové Znalosti")
    p_lang = doc.add_paragraph()
    p_lang.paragraph_format.space_before = Pt(3)
    p_lang.paragraph_format.space_after = Pt(6)
    
    langs = [
        ("Slovenský jazyk:", "C1 (Plynná úroveň / Plné profesionálne ovládanie)"),
        ("Anglický jazyk:", "B2 (Pracovná úroveň / Profesionálna komunikácia)"),
        ("Ukrajinský jazyk:", "Materinský jazyk (Native)"),
        ("Ruský jazyk:", "C2 (Plynulá úroveň)")
    ]
    for l_name, l_level in langs:
        r_ln = p_lang.add_run(f"{l_name} ")
        r_ln.font.bold = True; r_ln.font.size = Pt(9.5); r_ln.font.color.rgb = COLOR_NAVY
        r_lv = p_lang.add_run(f"{l_level}    •    ")
        r_lv.font.size = Pt(9.5); r_lv.font.color.rgb = COLOR_DARK

    doc.save("/Users/ruslana.buinytskawbpo.sk/CV_Web/Ruslana_Buinytska_CV_SK.docx")
    print("Created Ruslana_Buinytska_CV_SK.docx successfully!")

def build_cv_en():
    doc = docx.Document()
    
    # Page Setup (A4, 1.8 cm margins)
    section = doc.sections[0]
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)

    # 1. HEADER
    p_name = doc.add_paragraph()
    p_name.paragraph_format.space_before = Pt(0)
    p_name.paragraph_format.space_after = Pt(2)
    run_name = p_name.add_run("Ruslana Buinytska, MSc.")
    run_name.font.name = "Calibri"
    run_name.font.size = Pt(22)
    run_name.font.bold = True
    run_name.font.color.rgb = COLOR_NAVY

    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(6)
    run_title = p_title.add_run("IT Analyst (Systems & Business Analysis)  •  IT Project Manager")
    run_title.font.name = "Calibri"
    run_title.font.size = Pt(12)
    run_title.font.bold = True
    run_title.font.color.rgb = COLOR_DARK

    # Contact line 1: Location, Phone, Email
    p_contact1 = doc.add_paragraph()
    p_contact1.paragraph_format.space_before = Pt(0)
    p_contact1.paragraph_format.space_after = Pt(2)
    
    r = p_contact1.add_run("Bratislava, Slovakia  |  Phone: ")
    r.font.size = Pt(9.5); r.font.color.rgb = COLOR_MUTED
    r_tel = p_contact1.add_run("+421 951 698 832")
    r_tel.font.size = Pt(9.5); r_tel.font.bold = True; r_tel.font.color.rgb = COLOR_DARK
    
    r2 = p_contact1.add_run("  |  Email: ")
    r2.font.size = Pt(9.5); r2.font.color.rgb = COLOR_MUTED
    add_hyperlink(p_contact1, "mailto:rusbuin.oct@gmail.com", "rusbuin.oct@gmail.com", HEX_NAVY)

    # Contact line 2: Web Portfolio, LinkedIn, GitHub
    p_contact2 = doc.add_paragraph()
    p_contact2.paragraph_format.space_before = Pt(0)
    p_contact2.paragraph_format.space_after = Pt(10)
    
    r3 = p_contact2.add_run("Web Portfolio: ")
    r3.font.size = Pt(9.5); r3.font.bold = True; r3.font.color.rgb = COLOR_NAVY
    add_hyperlink(p_contact2, "https://rbuinytska.github.io", "rbuinytska.github.io", HEX_NAVY, bold=True)
    
    r4 = p_contact2.add_run("  |  LinkedIn: ")
    r4.font.size = Pt(9.5); r4.font.color.rgb = COLOR_MUTED
    add_hyperlink(p_contact2, "https://linkedin.com/in/ruslana-buinytska-188602395", "linkedin.com/in/ruslana-buinytska", HEX_NAVY)
    
    r5 = p_contact2.add_run("  |  GitHub: ")
    r5.font.size = Pt(9.5); r5.font.color.rgb = COLOR_MUTED
    add_hyperlink(p_contact2, "https://github.com/rbuinytska", "github.com/rbuinytska", HEX_NAVY)

    # 2. PROFILE
    add_section_heading(doc, "Professional Summary")
    p_profile = doc.add_paragraph()
    p_profile.paragraph_format.space_before = Pt(4)
    p_profile.paragraph_format.space_after = Pt(8)
    p_profile.paragraph_format.line_spacing = 1.15
    run_prof = p_profile.add_run(
        "Information Management engineer (EUBA) specializing in IT systems analysis, IT project governance, "
        "and data architecture. Grounded in practical software engineering experience, I understand software "
        "architecture from the bottom up, enabling me to author gap-free functional and technical specifications, "
        "design integration interfaces (REST APIs, JSON schemas), coordinate agile delivery teams (IPMA certified, "
        "SCRUM), and act as a credible, equal partner to both software developers and executive management."
    )
    run_prof.font.size = Pt(10)
    run_prof.font.color.rgb = COLOR_DARK

    # 3. CORE COMPETENCIES
    add_section_heading(doc, "Core Competencies & Methodologies")
    
    skills_data = [
        ("IT Analysis & Systems Modeling:", "UML (Use Case, Activity, Sequence diagrams), BPMN workflows, Enterprise Architect (EA), Draw.io, functional & technical specifications, requirements decomposition, API interface contracts (REST, JSON), User Stories, and Acceptance Criteria."),
        ("Project Management & Agile:", "IPMA Project Management Fundamentals international certification (2025), SCRUM & Kanban frameworks, sprint backlog prioritization, risk mitigation, milestone tracking, PRINCE2 & ITIL principles."),
        ("Data & BI Analytics:", "Relational databases, advanced Oracle SQL, data modeling (Star Schema), ETL data pipeline transformation, interactive dashboards in Power BI (DAX) and Tableau, Python data analysis."),
        ("Software Engineering & AI:", "REST API architecture, JSON/XML, Kotlin (KMP), Java, Android, Git, AI-Driven Development (LLM integration, AI agent orchestration).")
    ]
    
    for cat_title, cat_desc in skills_data:
        p_sk = doc.add_paragraph(style='List Bullet')
        p_sk.paragraph_format.space_before = Pt(2)
        p_sk.paragraph_format.space_after = Pt(2)
        p_sk.paragraph_format.line_spacing = 1.1
        r_t = p_sk.add_run(cat_title + " ")
        r_t.font.bold = True
        r_t.font.size = Pt(9.5)
        r_t.font.color.rgb = COLOR_NAVY
        r_d = p_sk.add_run(cat_desc)
        r_d.font.size = Pt(9.5)
        r_d.font.color.rgb = COLOR_DARK

    # 4. EXPERIENCE
    add_section_heading(doc, "Work Experience")

    # WBPO
    p_wbpo_h = doc.add_paragraph()
    p_wbpo_h.paragraph_format.space_before = Pt(6)
    p_wbpo_h.paragraph_format.space_after = Pt(2)
    p_wbpo_h.paragraph_format.keep_with_next = True
    r_role = p_wbpo_h.add_run("IT Analysis & Software Development")
    r_role.font.bold = True; r_role.font.size = Pt(11); r_role.font.color.rgb = COLOR_DARK
    r_comp = p_wbpo_h.add_run("  |  WBPO s.r.o. (Bratislava, Slovakia)")
    r_comp.font.bold = True; r_comp.font.size = Pt(10.5); r_comp.font.color.rgb = COLOR_NAVY
    
    p_wbpo_d = doc.add_paragraph()
    p_wbpo_d.paragraph_format.space_before = Pt(0)
    p_wbpo_d.paragraph_format.space_after = Pt(3)
    p_wbpo_d.paragraph_format.keep_with_next = True
    r_date = p_wbpo_d.add_run("2025 – Present  •  Current Position")
    r_date.font.size = Pt(9); r_date.font.italic = True; r_date.font.color.rgb = COLOR_MUTED

    wbpo_bullets = [
        "Independently broke down business stakeholder requirements into robust technical analysis and functional application specifications for software engineering teams.",
        "Engineered integration interface contracts (REST API definitions, JSON data schemas) with close coordination alongside backend architects and developers.",
        "Formulated comprehensive acceptance criteria, defined test scenarios, identified and resolved edge cases and exception handling flows.",
        "Collaborated within an agile SCRUM team, participating in sprint plannings, refinements, and bridging communication gaps between business and technology."
    ]
    for b in wbpo_bullets:
        p_b = doc.add_paragraph(style='List Bullet')
        p_b.paragraph_format.space_before = Pt(1)
        p_b.paragraph_format.space_after = Pt(2)
        p_b.paragraph_format.line_spacing = 1.1
        r_b = p_b.add_run(b)
        r_b.font.size = Pt(9.5)
        r_b.font.color.rgb = COLOR_DARK

    p_wbpo_t = doc.add_paragraph()
    p_wbpo_t.paragraph_format.space_before = Pt(2)
    p_wbpo_t.paragraph_format.space_after = Pt(6)
    r_tl = p_wbpo_t.add_run("Technologies & Tools: ")
    r_tl.font.bold = True; r_tl.font.size = Pt(9); r_tl.font.color.rgb = COLOR_MUTED
    r_tv = p_wbpo_t.add_run("IT Systems Analysis, REST API, JSON, UML, Agile / SCRUM, Technical Specifications, Kotlin (KMP), Git")
    r_tv.font.size = Pt(9); r_tv.font.color.rgb = COLOR_MUTED

    # 5. PROJECTS
    add_section_heading(doc, "Selected Projects & Portfolio")

    # ESG
    p_esg_h = doc.add_paragraph()
    p_esg_h.paragraph_format.space_before = Pt(6)
    p_esg_h.paragraph_format.space_after = Pt(2)
    p_esg_h.paragraph_format.keep_with_next = True
    r_p1 = p_esg_h.add_run("ESG Data Analysis & Visualization  ")
    r_p1.font.bold = True; r_p1.font.size = Pt(10.5); r_p1.font.color.rgb = COLOR_DARK
    r_p1s = p_esg_h.add_run("|  Master's Thesis (Ing.), EUBA  |  ")
    r_p1s.font.size = Pt(9.5); r_p1s.font.color.rgb = COLOR_MUTED
    add_hyperlink(p_esg_h, "https://github.com/RusBuin/esg-analysis-powerbi", "github.com/RusBuin/esg-analysis-powerbi", HEX_NAVY)

    esg_bullets = [
        "End-to-end data analytics project: automated collection, validation, cleaning, and ETL transformation of ESG indicators.",
        "Designed dimensional data models (Star Schema) and engineered advanced business calculations and KPI measures using DAX.",
        "Developed interactive corporate reporting dashboards in Power BI to drive executive management decision-making."
    ]
    for b in esg_bullets:
        p_b = doc.add_paragraph(style='List Bullet')
        p_b.paragraph_format.space_before = Pt(1)
        p_b.paragraph_format.space_after = Pt(2)
        r_b = p_b.add_run(b)
        r_b.font.size = Pt(9.5); r_b.font.color.rgb = COLOR_DARK

    # AI / ML
    p_ml_h = doc.add_paragraph()
    p_ml_h.paragraph_format.space_before = Pt(4)
    p_ml_h.paragraph_format.space_after = Pt(2)
    p_ml_h.paragraph_format.keep_with_next = True
    r_p2 = p_ml_h.add_run("AI Model Training & Data Analysis  ")
    r_p2.font.bold = True; r_p2.font.size = Pt(10.5); r_p2.font.color.rgb = COLOR_DARK
    r_p2s = p_ml_h.add_run("|  Bachelor's Thesis (Bc.), TUKE  |  ")
    r_p2s.font.size = Pt(9.5); r_p2s.font.color.rgb = COLOR_MUTED
    add_hyperlink(p_ml_h, "https://github.com/RusBuin/bakalarska_praca.git", "github.com/RusBuin/bakalarska_praca", HEX_NAVY)

    ml_bullets = [
        "Complex dataset exploratory analysis, statistical data preprocessing, and predictive machine learning model training in Python.",
        "Algorithm parameter tuning, predictive performance benchmarking, and analytical challenge resolution."
    ]
    for b in ml_bullets:
        p_b = doc.add_paragraph(style='List Bullet')
        p_b.paragraph_format.space_before = Pt(1)
        p_b.paragraph_format.space_after = Pt(2)
        r_b = p_b.add_run(b)
        r_b.font.size = Pt(9.5); r_b.font.color.rgb = COLOR_DARK

    # 6. EDUCATION
    add_section_heading(doc, "Education")
    
    edu_list = [
        ("2026 – Present", "PhD. in Data Science in Economics", "University of Economics in Bratislava (EUBA)", "Doctoral research focused on advanced data science methodologies and AI integration in economic systems."),
        ("2024 – 2026", "MSc. (Ing.) in Information Management", "University of Economics in Bratislava (EUBA)", "Graduate studies: IT process management, enterprise information systems, systems analysis, and project management."),
        ("2022 – 2024", "BSc. (Bc.) in Business Informatics", "Technical University of Košice (TUKE)", "Undergraduate studies: business informatics, software engineering, relational databases, and data structures.")
    ]
    for yr, deg, schl, desc in edu_list:
        p_e = doc.add_paragraph()
        p_e.paragraph_format.space_before = Pt(3)
        p_e.paragraph_format.space_after = Pt(2)
        r_deg = p_e.add_run(f"{deg}  |  {schl}")
        r_deg.font.bold = True; r_deg.font.size = Pt(10); r_deg.font.color.rgb = COLOR_DARK
        r_yr = p_e.add_run(f"  ({yr})\n")
        r_yr.font.size = Pt(9); r_yr.font.color.rgb = COLOR_MUTED
        r_dc = p_e.add_run(desc)
        r_dc.font.size = Pt(9.5); r_dc.font.color.rgb = COLOR_MUTED

    # 7. CERTIFICATIONS
    add_section_heading(doc, "Certifications & Professional Courses")
    certs = [
        ("2025", "IPMA Project Management Fundamentals Certification", "IPMA Slovakia – International credential in project planning, execution, and control methodologies."),
        ("2025", "Sustainable Project Management #HACKATHON2025", "Practical challenge resolution applying agile project governance and cross-functional team leadership."),
        ("2022", "Agile Workshop „LEGO SCRUM & TUKE“", "Technical University of Košice – Hands-on agile ceremonies, sprint simulations, and SCRUM framework."),
        ("2022", "„Skills for Success“ Leadership Program (AmCham Slovakia)", "American Chamber of Commerce – Intensive managerial skill development, negotiation, and corporate communications.")
    ]
    for yr, title, desc in certs:
        p_c = doc.add_paragraph(style='List Bullet')
        p_c.paragraph_format.space_before = Pt(1)
        p_c.paragraph_format.space_after = Pt(2)
        r_ct = p_c.add_run(f"{title} ({yr}): ")
        r_ct.font.bold = True; r_ct.font.size = Pt(9.5); r_ct.font.color.rgb = COLOR_DARK
        r_cd = p_c.add_run(desc)
        r_cd.font.size = Pt(9.5); r_cd.font.color.rgb = COLOR_MUTED

    # 8. LANGUAGES
    add_section_heading(doc, "Languages")
    p_lang = doc.add_paragraph()
    p_lang.paragraph_format.space_before = Pt(3)
    p_lang.paragraph_format.space_after = Pt(6)
    
    langs = [
        ("Slovak:", "C1 (Full Professional Proficiency)"),
        ("English:", "B2 (Professional Working Proficiency)"),
        ("Ukrainian:", "Native"),
        ("Russian:", "C2 (Bilingual / Fluent)")
    ]
    for l_name, l_level in langs:
        r_ln = p_lang.add_run(f"{l_name} ")
        r_ln.font.bold = True; r_ln.font.size = Pt(9.5); r_ln.font.color.rgb = COLOR_NAVY
        r_lv = p_lang.add_run(f"{l_level}    •    ")
        r_lv.font.size = Pt(9.5); r_lv.font.color.rgb = COLOR_DARK

    doc.save("/Users/ruslana.buinytskawbpo.sk/CV_Web/Ruslana_Buinytska_CV_EN.docx")
    print("Created Ruslana_Buinytska_CV_EN.docx successfully!")

if __name__ == "__main__":
    build_cv_sk()
    build_cv_en()
