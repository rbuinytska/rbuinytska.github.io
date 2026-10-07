import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn

COLOR_NAVY = RGBColor(30, 58, 138)       # #1E3A8A
COLOR_DARK = RGBColor(15, 23, 42)        # #0F172A
COLOR_MUTED = RGBColor(71, 85, 105)      # #475569
HEX_BLUE = "2563EB"
HEX_NAVY = "1E3A8A"

def add_hyperlink(paragraph, url, text, color_hex="2563EB", bold=True, size_pt=14):
    part = paragraph.part
    r_id = part.relate_to(url, docx.opc.constants.RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    hyperlink = OxmlElement('w:hyperlink')
    hyperlink.set(qn('r:id'), r_id)
    new_run = OxmlElement('w:r')
    rPr = OxmlElement('w:rPr')
    if color_hex:
        c = OxmlElement('w:color')
        c.set(qn('w:val'), color_hex)
        rPr.append(c)
    u = OxmlElement('w:u')
    u.set(qn('w:val'), 'single')
    rPr.append(u)
    if bold:
        b = OxmlElement('w:b')
        rPr.append(b)
    if size_pt:
        sz = OxmlElement('w:sz')
        sz.set(qn('w:val'), str(int(size_pt * 2)))
        rPr.append(sz)
    new_run.append(rPr)
    new_run.text = text
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)

def make_doc(lang="sk"):
    doc = docx.Document()
    
    # Page setup
    section = doc.sections[0]
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.top_margin = Inches(1.5)
    section.bottom_margin = Inches(1.5)
    section.left_margin = Inches(1.2)
    section.right_margin = Inches(1.2)

    # Name
    p_name = doc.add_paragraph()
    p_name.paragraph_format.space_before = Pt(0)
    p_name.paragraph_format.space_after = Pt(4)
    run_name = p_name.add_run("Ing. Ruslana Buinytska" if lang == "sk" else "Ruslana Buinytska, MSc.")
    run_name.font.name = "Calibri"
    run_name.font.size = Pt(26)
    run_name.font.bold = True
    run_name.font.color.rgb = COLOR_NAVY

    # Subtitle / Role
    p_role = doc.add_paragraph()
    p_role.paragraph_format.space_before = Pt(0)
    p_role.paragraph_format.space_after = Pt(12)
    run_role = p_role.add_run("IT Analytička  •  IT Project Manager" if lang == "sk" else "IT Analyst  •  IT Project Manager")
    run_role.font.name = "Calibri"
    run_role.font.size = Pt(13)
    run_role.font.bold = True
    run_role.font.color.rgb = COLOR_DARK

    # Hairline divider
    p_div = doc.add_paragraph()
    p_div.paragraph_format.space_before = Pt(0)
    p_div.paragraph_format.space_after = Pt(20)
    pBdr = parse_xml(r'<w:pBdr xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
                     r'<w:bottom w:val="single" w:sz="8" w:space="1" w:color="1E3A8A"/>'
                     r'</w:pBdr>')
    p_div._p.get_or_add_pPr().append(pBdr)

    # Short Bio Text
    p_text = doc.add_paragraph()
    p_text.paragraph_format.space_before = Pt(0)
    p_text.paragraph_format.space_after = Pt(24)
    p_text.paragraph_format.line_spacing = 1.25

    if lang == "sk":
        bio_text = (
            "Som inžinierka Informačného manažmentu (EUBA, certifikácia IPMA) so zameraním na IT systémovú analýzu, "
            "riadenie IT projektov a dáta. Prepájam biznisové požiadavky s technickou realitou. Vďaka praktickej skúsenosti "
            "zo softvérového vývoja rozumiem architektúre systémov do hĺbky, čo mi umožňuje tvoriť precízne funkčné a technické "
            "špecifikácie bez logických medzier, viesť agilné tímy a byť rovnocenným partnerom pre vývojárov aj manažment."
        )
    else:
        bio_text = (
            "I am an Information Management engineer (EUBA, IPMA certified) specializing in IT systems analysis, "
            "IT project governance, and data architecture. I bridge business requirements with software engineering realities. "
            "With a hands-on foundation in software development, I understand systems architecture from the ground up, allowing me "
            "to author gap-free technical specifications, coordinate agile teams, and act as an equal partner to software engineers and leadership."
        )
    r_bio = p_text.add_run(bio_text)
    r_bio.font.name = "Calibri"
    r_bio.font.size = Pt(11)
    r_bio.font.color.rgb = COLOR_DARK

    # Link introduction text
    p_link_intro = doc.add_paragraph()
    p_link_intro.paragraph_format.space_before = Pt(0)
    p_link_intro.paragraph_format.space_after = Pt(8)
    r_li = p_link_intro.add_run(
        "Kompletný interaktívny životopis, projekty a špecifikácie nájdete na:" if lang == "sk"
        else "Complete interactive resume, projects, and specifications can be found at:"
    )
    r_li.font.name = "Calibri"
    r_li.font.size = Pt(11)
    r_li.font.color.rgb = COLOR_MUTED

    # Prominent Link
    p_link = doc.add_paragraph()
    p_link.paragraph_format.space_before = Pt(0)
    p_link.paragraph_format.space_after = Pt(30)
    add_hyperlink(p_link, "https://rbuinytska.github.io", "https://rbuinytska.github.io", HEX_BLUE, bold=True, size_pt=16)

    # Hairline divider
    p_div2 = doc.add_paragraph()
    p_div2.paragraph_format.space_before = Pt(0)
    p_div2.paragraph_format.space_after = Pt(14)
    pBdr2 = parse_xml(r'<w:pBdr xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
                      r'<w:bottom w:val="single" w:sz="6" w:space="1" w:color="CBD5E1"/>'
                      r'</w:pBdr>')
    p_div2._p.get_or_add_pPr().append(pBdr2)

    # Contacts
    p_cnt = doc.add_paragraph()
    p_cnt.paragraph_format.space_before = Pt(0)
    p_cnt.paragraph_format.space_after = Pt(4)
    p_cnt.paragraph_format.line_spacing = 1.2
    
    r_c = p_cnt.add_run("Bratislava, Slovensko  |  Tel: +421 951 698 832  |  Email: rusbuin.oct@gmail.com\nLinkedIn: linkedin.com/in/ruslana-buinytska  |  GitHub: github.com/rbuinytska" if lang == "sk"
                        else "Bratislava, Slovakia  |  Phone: +421 951 698 832  |  Email: rusbuin.oct@gmail.com\nLinkedIn: linkedin.com/in/ruslana-buinytska  |  GitHub: github.com/rbuinytska")
    r_c.font.name = "Calibri"
    r_c.font.size = Pt(9.5)
    r_c.font.color.rgb = COLOR_MUTED

    filename = "/Users/ruslana.buinytskawbpo.sk/CV_Web/Ruslana_Buinytska_CV.docx" if lang == "sk" else "/Users/ruslana.buinytskawbpo.sk/CV_Web/Ruslana_Buinytska_CV_EN.docx"
    doc.save(filename)
    print("Saved:", filename)

make_doc("sk")
make_doc("en")
