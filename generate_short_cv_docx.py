import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn

# Colors
COLOR_NAVY = RGBColor(30, 58, 138)       # #1E3A8A
COLOR_DARK = RGBColor(15, 23, 42)        # #0F172A
COLOR_MUTED = RGBColor(71, 85, 105)      # #475569
COLOR_BLUE_LINK = RGBColor(37, 99, 235)  # #2563EB
HEX_NAVY = "1E3A8A"
HEX_BLUE = "2563EB"
HEX_BG_BOX = "F1F5F9"                    # Soft slate box

def add_hyperlink(paragraph, url, text, color_hex="1E3A8A", bold=False, size_pt=None):
    part = paragraph.part
    r_id = part.relate_to(url, docx.opc.constants.RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    hyperlink = OxmlElement('w:hyperlink')
    hyperlink.set(qn('r:id'), r_id)
    new_run = OxmlElement('w:r')
    rPr = OxmlElement('w:rPr')
    if color_hex:
        c = OxmlElement('w:color')
        c.set(qn('w:val'), color_hex)
        rPr.append(c)
    u = OxmlElement('w:u')
    u.set(qn('w:val'), 'single')
    rPr.append(u)
    if bold:
        b = OxmlElement('w:b')
        rPr.append(b)
    if size_pt:
        sz = OxmlElement('w:sz')
        sz.set(qn('w:val'), str(int(size_pt * 2)))
        rPr.append(sz)
    new_run.append(rPr)
    new_run.text = text
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)

def create_short_cv_sk():
    doc = docx.Document()
    
    # Page setup (A4, 2 cm margins)
    section = doc.sections[0]
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)

    # 1. NAME & TITLE
    p_name = doc.add_paragraph()
    p_name.paragraph_format.space_before = Pt(0)
    p_name.paragraph_format.space_after = Pt(2)
    p_name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_name = p_name.add_run("Ing. Ruslana Buinytska")
    r_name.font.name = "Calibri"
    r_name.font.size = Pt(24)
    r_name.font.bold = True
    r_name.font.color.rgb = COLOR_NAVY

    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(8)
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_title = p_title.add_run("IT Analytička  •  IT Project Manager")
    r_title.font.name = "Calibri"
    r_title.font.size = Pt(13)
    r_title.font.bold = True
    r_title.font.color.rgb = COLOR_DARK

    # CONTACT INFO
    p_contact = doc.add_paragraph()
    p_contact.paragraph_format.space_before = Pt(0)
    p_contact.paragraph_format.space_after = Pt(16)
    p_contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    r_c1 = p_contact.add_run("Bratislava, Slovensko  •  Tel: +421 951 698 832  •  Email: ")
    r_c1.font.size = Pt(9.5); r_c1.font.color.rgb = COLOR_MUTED
    add_hyperlink(p_contact, "mailto:rusbuin.oct@gmail.com", "rusbuin.oct@gmail.com", HEX_NAVY)
    
    r_c2 = p_contact.add_run("\nLinkedIn: ")
    r_c2.font.size = Pt(9.5); r_c2.font.color.rgb = COLOR_MUTED
    add_hyperlink(p_contact, "https://linkedin.com/in/ruslana-buinytska-188602395", "linkedin.com/in/ruslana-buinytska", HEX_NAVY)
    r_c3 = p_contact.add_run("  •  GitHub: ")
    r_c3.font.size = Pt(9.5); r_c3.font.color.rgb = COLOR_MUTED
    add_hyperlink(p_contact, "https://github.com/rbuinytska", "github.com/rbuinytska", HEX_NAVY)

    # 2. PROMINENT CALLOUT BOX WITH LINK TO PORTFOLIO
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    cell.width = Inches(6.5)
    
    # Background color and border for the callout box
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(r'<w:shd xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:val="clear" w:color="auto" w:fill="F1F5F9"/>')
    tcPr.append(shd)
    borders = parse_xml(r'<w:tcBorders xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
                        r'<w:top w:val="single" w:sz="12" w:space="0" w:color="1E3A8A"/>'
                        r'<w:left w:val="single" w:sz="12" w:space="0" w:color="1E3A8A"/>'
                        r'<w:bottom w:val="single" w:sz="12" w:space="0" w:color="1E3A8A"/>'
                        r'<w:right w:val="single" w:sz="12" w:space="0" w:color="1E3A8A"/>'
                        r'</w:tcBorders>')
    tcPr.append(borders)
    
    # Inside Box: Heading
    p_box1 = cell.paragraphs[0]
    p_box1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_box1.paragraph_format.space_before = Pt(8)
    p_box1.paragraph_format.space_after = Pt(4)
    r_box_h = p_box1.add_run("INTERAKTÍVNE PORTFÓLIO & KOMPLETNÝ ŽIVOTOPIS")
    r_box_h.font.bold = True
    r_box_h.font.size = Pt(11)
    r_box_h.font.color.rgb = COLOR_NAVY
    
    # Inside Box: The Prominent Clickable Link
    p_box2 = cell.add_paragraph()
    p_box2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_box2.paragraph_format.space_before = Pt(0)
    p_box2.paragraph_format.space_after = Pt(4)
    add_hyperlink(p_box2, "https://rbuinytska.github.io", "https://rbuinytska.github.io", HEX_BLUE, bold=True, size_pt=14)
    
    # Inside Box: Subtitle note
    p_box3 = cell.add_paragraph()
    p_box3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_box3.paragraph_format.space_before = Pt(0)
    p_box3.paragraph_format.space_after = Pt(8)
    r_box_sub = p_box3.add_run("Kliknite na odkaz vyššie pre zobrazenie interaktívnych projektov, systémových špecifikácií a certifikácií.")
    r_box_sub.font.italic = True
    r_box_sub.font.size = Pt(9.5)
    r_box_sub.font.color.rgb = COLOR_MUTED

    # 3. SHORT BIO / PROFIL
    p_sh = doc.add_paragraph()
    p_sh.paragraph_format.space_before = Pt(18)
    p_sh.paragraph_format.space_after = Pt(4)
    r_sh = p_sh.add_run("STRUČNÝ PROFIL")
    r_sh.font.bold = True
    r_sh.font.size = Pt(11)
    r_sh.font.color.rgb = COLOR_NAVY
    pBdr = parse_xml(r'<w:pBdr xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
                     r'<w:bottom w:val="single" w:sz="6" w:space="2" w:color="1E3A8A"/>'
                     r'</w:pBdr>')
    p_sh._p.get_or_add_pPr().append(pBdr)

    p_bio = doc.add_paragraph()
    p_bio.paragraph_format.space_before = Pt(4)
    p_bio.paragraph_format.space_after = Pt(12)
    p_bio.paragraph_format.line_spacing = 1.15
    r_b = p_bio.add_run(
        "Som inžinierka Informačného manažmentu (EUBA) so zameraním na IT analýzu, riadenie IT projektov a dátovú architektúru. "
        "Vďaka praktickej skúsenosti zo softvérového vývoja rozumiem architektúre systémov do hĺbky, čo mi umožňuje tvoriť "
        "precízne funkčné a technické špecifikácie bez logických medzier, navrhovať REST API a JSON schémy, viesť agilné tímy "
        "(certifikácia IPMA, SCRUM) a byť rovnocenným partnerom pre vývojárov aj manažment."
    )
    r_b.font.size = Pt(10)
    r_b.font.color.rgb = COLOR_DARK

    # 4. KĽÚČOVÉ PILIERE
    p_kh = doc.add_paragraph()
    p_kh.paragraph_format.space_before = Pt(10)
    p_kh.paragraph_format.space_after = Pt(4)
    r_kh = p_kh.add_run("KĽÚČOVÉ OBLASTI PÔSOBENIA")
    r_kh.font.bold = True
    r_kh.font.size = Pt(11)
    r_kh.font.color.rgb = COLOR_NAVY
    pBdr2 = parse_xml(r'<w:pBdr xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
                      r'<w:bottom w:val="single" w:sz="6" w:space="2" w:color="1E3A8A"/>'
                      r'</w:pBdr>')
    p_kh._p.get_or_add_pPr().append(pBdr2)

    pillars = [
        ("IT Systémová & Biznis Analýza: ", "Modelovanie v UML a BPMN (Enterprise Architect), dekompozícia požiadaviek stakeholderov, funkčné špecifikácie, návrh API kontraktov (REST, JSON) a akceptačné testovanie."),
        ("IT Projektový Manažment: ", "Medzinárodná certifikácia IPMA Project Management Fundamentals (2025), agilné rámce SCRUM & Kanban, prioritizácia sprintov, riadenie rizík a koordinácia dodávky."),
        ("Dáta, SQL & Power BI: ", "Relačné databázy, Oracle SQL, dátové modelovanie (Star Schema), ETL toky dát, interaktívne manažérske dashboardy v Power BI (DAX) a Tableau."),
        ("Technický Vývoj & AI: ", "Prax z vývoja (Kotlin/KMP, Java, Git), návrh aplikačnej logiky, integrácia LLM modelov a AI agentov.")
    ]
    for p_title_text, p_desc_text in pillars:
        p_pil = doc.add_paragraph(style='List Bullet')
        p_pil.paragraph_format.space_before = Pt(2)
        p_pil.paragraph_format.space_after = Pt(2)
        r_pt = p_pil.add_run(p_title_text)
        r_pt.font.bold = True; r_pt.font.size = Pt(9.5); r_pt.font.color.rgb = COLOR_NAVY
        r_pd = p_pil.add_run(p_desc_text)
        r_pd.font.size = Pt(9.5); r_pd.font.color.rgb = COLOR_DARK

    # 5. VZDELANIE & KVALIFIKÁCIA
    p_eh = doc.add_paragraph()
    p_eh.paragraph_format.space_before = Pt(12)
    p_eh.paragraph_format.space_after = Pt(4)
    r_eh = p_eh.add_run("VZDELANIE & CERTIFIKÁCIA")
    r_eh.font.bold = True
    r_eh.font.size = Pt(11)
    r_eh.font.color.rgb = COLOR_NAVY
    pBdr3 = parse_xml(r'<w:pBdr xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
                      r'<w:bottom w:val="single" w:sz="6" w:space="2" w:color="1E3A8A"/>'
                      r'</w:pBdr>')
    p_eh._p.get_or_add_pPr().append(pBdr3)

    edu_items = [
        ("IPMA Project Management Fundamentals", "Medzinárodná certifikácia projektového manažmentu (2025)"),
        ("Informačný manažment (Ing.)", "Ekonomická univerzita v Bratislave (2024 – 2026)"),
        ("Hospodárska informatika (Bc.)", "Technická univerzita v Košiciach (2022 – 2024)")
    ]
    for title, sub in edu_items:
        p_e = doc.add_paragraph(style='List Bullet')
        p_e.paragraph_format.space_before = Pt(1)
        p_e.paragraph_format.space_after = Pt(2)
        r_et = p_e.add_run(title + " — ")
        r_et.font.bold = True; r_et.font.size = Pt(9.5); r_et.font.color.rgb = COLOR_DARK
        r_es = p_e.add_run(sub)
        r_es.font.size = Pt(9.5); r_es.font.color.rgb = COLOR_MUTED

    # SAVE
    doc.save("/Users/ruslana.buinytskawbpo.sk/CV_Web/Ruslana_Buinytska_CV.docx")
    print("Created Ruslana_Buinytska_CV.docx (Short 1-page version)")

def create_short_cv_en():
    doc = docx.Document()
    
    # Page setup
    section = doc.sections[0]
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)

    # 1. NAME & TITLE
    p_name = doc.add_paragraph()
    p_name.paragraph_format.space_before = Pt(0)
    p_name.paragraph_format.space_after = Pt(2)
    p_name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_name = p_name.add_run("Ruslana Buinytska, MSc.")
    r_name.font.name = "Calibri"
    r_name.font.size = Pt(24)
    r_name.font.bold = True
    r_name.font.color.rgb = COLOR_NAVY

    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(8)
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_title = p_title.add_run("IT Analyst  •  IT Project Manager")
    r_title.font.name = "Calibri"
    r_title.font.size = Pt(13)
    r_title.font.bold = True
    r_title.font.color.rgb = COLOR_DARK

    # CONTACT INFO
    p_contact = doc.add_paragraph()
    p_contact.paragraph_format.space_before = Pt(0)
    p_contact.paragraph_format.space_after = Pt(16)
    p_contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    r_c1 = p_contact.add_run("Bratislava, Slovakia  •  Phone: +421 951 698 832  •  Email: ")
    r_c1.font.size = Pt(9.5); r_c1.font.color.rgb = COLOR_MUTED
    add_hyperlink(p_contact, "mailto:rusbuin.oct@gmail.com", "rusbuin.oct@gmail.com", HEX_NAVY)
    
    r_c2 = p_contact.add_run("\nLinkedIn: ")
    r_c2.font.size = Pt(9.5); r_c2.font.color.rgb = COLOR_MUTED
    add_hyperlink(p_contact, "https://linkedin.com/in/ruslana-buinytska-188602395", "linkedin.com/in/ruslana-buinytska", HEX_NAVY)
    r_c3 = p_contact.add_run("  •  GitHub: ")
    r_c3.font.size = Pt(9.5); r_c3.font.color.rgb = COLOR_MUTED
    add_hyperlink(p_contact, "https://github.com/rbuinytska", "github.com/rbuinytska", HEX_NAVY)

    # 2. PROMINENT CALLOUT BOX
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    cell.width = Inches(6.5)
    
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(r'<w:shd xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:val="clear" w:color="auto" w:fill="F1F5F9"/>')
    tcPr.append(shd)
    borders = parse_xml(r'<w:tcBorders xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
                        r'<w:top w:val="single" w:sz="12" w:space="0" w:color="1E3A8A"/>'
                        r'<w:left w:val="single" w:sz="12" w:space="0" w:color="1E3A8A"/>'
                        r'<w:bottom w:val="single" w:sz="12" w:space="0" w:color="1E3A8A"/>'
                        r'<w:right w:val="single" w:sz="12" w:space="0" w:color="1E3A8A"/>'
                        r'</w:tcBorders>')
    tcPr.append(borders)
    
    p_box1 = cell.paragraphs[0]
    p_box1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_box1.paragraph_format.space_before = Pt(8)
    p_box1.paragraph_format.space_after = Pt(4)
    r_box_h = p_box1.add_run("INTERACTIVE PORTFOLIO & COMPLETE RESUME")
    r_box_h.font.bold = True
    r_box_h.font.size = Pt(11)
    r_box_h.font.color.rgb = COLOR_NAVY
    
    p_box2 = cell.add_paragraph()
    p_box2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_box2.paragraph_format.space_before = Pt(0)
    p_box2.paragraph_format.space_after = Pt(4)
    add_hyperlink(p_box2, "https://rbuinytska.github.io", "https://rbuinytska.github.io", HEX_BLUE, bold=True, size_pt=14)
    
    p_box3 = cell.add_paragraph()
    p_box3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_box3.paragraph_format.space_before = Pt(0)
    p_box3.paragraph_format.space_after = Pt(8)
    r_box_sub = p_box3.add_run("Click the link above to view interactive case studies, system specifications, and live projects.")
    r_box_sub.font.italic = True
    r_box_sub.font.size = Pt(9.5)
    r_box_sub.font.color.rgb = COLOR_MUTED

    # 3. SHORT PROFILE
    p_sh = doc.add_paragraph()
    p_sh.paragraph_format.space_before = Pt(18)
    p_sh.paragraph_format.space_after = Pt(4)
    r_sh = p_sh.add_run("EXECUTIVE PROFILE")
    r_sh.font.bold = True
    r_sh.font.size = Pt(11)
    r_sh.font.color.rgb = COLOR_NAVY
    pBdr = parse_xml(r'<w:pBdr xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
                     r'<w:bottom w:val="single" w:sz="6" w:space="2" w:color="1E3A8A"/>'
                     r'</w:pBdr>')
    p_sh._p.get_or_add_pPr().append(pBdr)

    p_bio = doc.add_paragraph()
    p_bio.paragraph_format.space_before = Pt(4)
    p_bio.paragraph_format.space_after = Pt(12)
    p_bio.paragraph_format.line_spacing = 1.15
    r_b = p_bio.add_run(
        "Information Management engineer (EUBA) specializing in IT systems analysis, project governance, and data architecture. "
        "With a hands-on background in software development, I understand systems architecture from the ground up, allowing me "
        "to author unambiguous functional and technical specifications, design REST APIs and JSON schemas, lead agile teams "
        "(IPMA certified, SCRUM), and act as a credible partner to software developers and business leaders alike."
    )
    r_b.font.size = Pt(10)
    r_b.font.color.rgb = COLOR_DARK

    # 4. CORE PILLARS
    p_kh = doc.add_paragraph()
    p_kh.paragraph_format.space_before = Pt(10)
    p_kh.paragraph_format.space_after = Pt(4)
    r_kh = p_kh.add_run("CORE COMPETENCIES")
    r_kh.font.bold = True
    r_kh.font.size = Pt(11)
    r_kh.font.color.rgb = COLOR_NAVY
    pBdr2 = parse_xml(r'<w:pBdr xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
                      r'<w:bottom w:val="single" w:sz="6" w:space="2" w:color="1E3A8A"/>'
                      r'</w:pBdr>')
    p_kh._p.get_or_add_pPr().append(pBdr2)

    pillars = [
        ("IT Systems & Business Analysis: ", "Process & systems modeling in UML/BPMN (Enterprise Architect), requirements decomposition, functional specs, API contract design (REST, JSON), and acceptance testing."),
        ("IT Project Governance: ", "IPMA Project Management Fundamentals international certification (2025), agile frameworks (SCRUM & Kanban), sprint prioritization, and delivery coordination."),
        ("Data, SQL & Business Intelligence: ", "Relational databases, Oracle SQL, dimensional data modeling (Star Schema), ETL data pipelines, and interactive executive dashboards in Power BI (DAX)."),
        ("Software Engineering & AI: ", "Direct development background (Kotlin/KMP, Java, Git), application architecture, and AI-driven development (LLM integration, AI agent orchestration).")
    ]
    for p_title_text, p_desc_text in pillars:
        p_pil = doc.add_paragraph(style='List Bullet')
        p_pil.paragraph_format.space_before = Pt(2)
        p_pil.paragraph_format.space_after = Pt(2)
        r_pt = p_pil.add_run(p_title_text)
        r_pt.font.bold = True; r_pt.font.size = Pt(9.5); r_pt.font.color.rgb = COLOR_NAVY
        r_pd = p_pil.add_run(p_desc_text)
        r_pd.font.size = Pt(9.5); r_pd.font.color.rgb = COLOR_DARK

    # 5. EDUCATION & CREDENTIALS
    p_eh = doc.add_paragraph()
    p_eh.paragraph_format.space_before = Pt(12)
    p_eh.paragraph_format.space_after = Pt(4)
    r_eh = p_eh.add_run("EDUCATION & CERTIFICATIONS")
    r_eh.font.bold = True
    r_eh.font.size = Pt(11)
    r_eh.font.color.rgb = COLOR_NAVY
    pBdr3 = parse_xml(r'<w:pBdr xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
                      r'<w:bottom w:val="single" w:sz="6" w:space="2" w:color="1E3A8A"/>'
                      r'</w:pBdr>')
    p_eh._p.get_or_add_pPr().append(pBdr3)

    edu_items = [
        ("IPMA Project Management Fundamentals", "International project management credential (2025)"),
        ("MSc. (Ing.) in Information Management", "University of Economics in Bratislava (2024 – 2026)"),
        ("BSc. (Bc.) in Business Informatics", "Technical University of Košice (2022 – 2024)")
    ]
    for title, sub in edu_items:
        p_e = doc.add_paragraph(style='List Bullet')
        p_e.paragraph_format.space_before = Pt(1)
        p_e.paragraph_format.space_after = Pt(2)
        r_et = p_e.add_run(title + " — ")
        r_et.font.bold = True; r_et.font.size = Pt(9.5); r_et.font.color.rgb = COLOR_DARK
        r_es = p_e.add_run(sub)
        r_es.font.size = Pt(9.5); r_es.font.color.rgb = COLOR_MUTED

    # SAVE
    doc.save("/Users/ruslana.buinytskawbpo.sk/CV_Web/Ruslana_Buinytska_CV_EN_Short.docx")
    print("Created Ruslana_Buinytska_CV_EN_Short.docx")

if __name__ == "__main__":
    create_short_cv_sk()
    create_short_cv_en()
